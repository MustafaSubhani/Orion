"""Ingest user files into kb/ only. Never touches data/frozen or bench/."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from harness.errors import PolicyDeniedError, SchemaInvalidError
from kb.store import (
    files_dir,
    iter_chunks,
    kb_table_names,
    root,
    tables_dir,
    upsert_resource,
    write_chunks,
)

MAX_BYTES = 15_000_000
MAX_TABLE_ROWS = 20_000
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 200

ALLOWED_EXT = frozenset(
    {
        ".csv",
        ".tsv",
        ".xlsx",
        ".xls",
        ".pdf",
        ".docx",
        ".txt",
        ".md",
        ".ics",
        ".json",
        ".pptx",
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
        ".gif",
    }
)
BANNED_EXT = frozenset({".py", ".exe", ".dll", ".env", ".bat", ".ps1", ".sh", ".cmd", ".msi"})
IMAGE_EXT = frozenset({".png", ".jpg", ".jpeg", ".webp", ".gif"})


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def safe_filename(name: str) -> str:
    base = Path(str(name).replace("\\", "/")).name
    base = re.sub(r"[^\w.\- ]+", "_", base).strip() or "upload.bin"
    if base in {".", ".."}:
        base = "upload.bin"
    return base[:180]


def _ext(name: str) -> str:
    return Path(name).suffix.lower()


def _slug(name: str) -> str:
    stem = Path(name).stem.lower()
    slug = re.sub(r"[^a-z0-9]+", "_", stem).strip("_")[:40]
    return slug or "file"


def _unique_table(slug: str) -> str:
    base = f"kb_{slug}"
    existing = set(kb_table_names())
    if base not in existing:
        return base
    for i in range(2, 100):
        candidate = f"{base}_{i}"
        if candidate not in existing:
            return candidate
    return f"{base}_{uuid.uuid4().hex[:6]}"


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = (text or "").replace("\x00", "")
    if not text.strip():
        return []
    chunks: list[str] = []
    i = 0
    n = len(text)
    step = max(1, size - overlap)
    while i < n:
        chunks.append(text[i : i + size])
        i += step
    return chunks


def _try_index(chunks: list[dict[str, Any]]) -> bool:
    if not chunks:
        return False
    try:
        from harness.config import COLLECTION_KB
        from retrieval.embeddings import get_embeddings_batch
        from retrieval.vector_store import get_or_create_collection

        col = get_or_create_collection(COLLECTION_KB)
        texts = [c["text"] for c in chunks]
        ids = [c["id"] for c in chunks]
        metas = [
            {
                "source": "kb",
                "filename": c.get("filename", ""),
                "resource_id": c.get("resource_id", ""),
            }
            for c in chunks
        ]
        embs = get_embeddings_batch(texts)
        col.upsert(ids=ids, documents=texts, embeddings=embs, metadatas=metas)
        return True
    except Exception:
        return False


def reindex_all() -> dict[str, Any]:
    """Rebuild the vector index from stored chunks. Lexical search keeps working if indexing is unavailable."""
    from kb.store import iter_chunks

    chunks = iter_chunks()
    indexed = _try_index(chunks) if chunks else False
    return {"chunks": len(chunks), "indexed": indexed}


def lexical_search(query: str, k: int = 8) -> list[dict[str, Any]]:
    terms = [t.lower() for t in re.findall(r"[a-zA-Z0-9_]+", query or "") if len(t) > 1]
    scored: list[tuple[int, dict[str, Any]]] = []
    for chunk in iter_chunks():
        blob = str(chunk.get("text") or "").lower()
        score = sum(blob.count(t) for t in terms) if terms else (1 if blob else 0)
        if not terms and blob:
            score = 1
        if score > 0:
            scored.append((score, chunk))
    scored.sort(key=lambda x: x[0], reverse=True)
    return [c for _, c in scored[: max(1, min(k, 20))]]


def search_kb(query: str, k: int = 8) -> list[dict[str, Any]]:
    k = max(1, min(int(k or 8), 20))
    q = (query or "").strip()
    if not q:
        raise SchemaInvalidError("kb_search.query is required")
    try:
        from harness.config import COLLECTION_KB
        from retrieval.vector_store import query_collection_diverse

        raw = query_collection_diverse(COLLECTION_KB, q, top_k=k)
        docs = (raw.get("documents") or [[]])[0]
        metas = (raw.get("metadatas") or [[]])[0]
        dists = (raw.get("distances") or [[]])[0]
        chunks = []
        for i, doc in enumerate(docs):
            meta = metas[i] if i < len(metas) else {}
            chunks.append(
                {
                    "chunk_id": meta.get("id") or f"kb:{i}",
                    "text": doc,
                    "metadata": meta,
                    "distance": dists[i] if i < len(dists) else None,
                }
            )
        if chunks:
            return chunks
    except Exception:
        pass
    hits = lexical_search(q, k)
    return [
        {
            "chunk_id": h.get("id"),
            "text": h.get("text"),
            "metadata": {"filename": h.get("filename"), "resource_id": h.get("resource_id")},
            "distance": None,
        }
        for h in hits
    ]


def _extract_text_pdf(data: bytes) -> str:
    try:
        import fitz

        doc = fitz.open(stream=data, filetype="pdf")
        try:
            return "\n".join(page.get_text() for page in doc)
        finally:
            doc.close()
    except Exception:
        pass
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception:
        try:
            from PyPDF2 import PdfReader as LegacyReader

            reader = LegacyReader(io.BytesIO(data))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception:
            return ""


def _extract_text_docx(data: bytes) -> str:
    try:
        import docx

        doc = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception:
        return ""


def _extract_text_pptx(data: bytes) -> str:
    try:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(data))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    parts.append(shape.text)
        return "\n".join(parts)
    except Exception:
        return ""


def _parse_ics(text: str) -> list[dict[str, str]]:
    events: list[dict[str, str]] = []
    current: dict[str, str] = {}
    key = ""
    for raw in text.splitlines():
        line = raw.rstrip("\r")
        if line.startswith(" ") and key:
            current[key] = current.get(key, "") + line[1:]
            continue
        if line == "BEGIN:VEVENT":
            current = {}
            continue
        if line == "END:VEVENT":
            if current:
                events.append(
                    {
                        "event_id": current.get("UID", ""),
                        "title": current.get("SUMMARY", ""),
                        "start_time": current.get("DTSTART", ""),
                        "end_time": current.get("DTEND", ""),
                        "organizer": current.get("ORGANIZER", ""),
                        "attendees": current.get("ATTENDEE", ""),
                        "location": current.get("LOCATION", ""),
                    }
                )
            current = {}
            continue
        if ":" in line:
            k, _, v = line.partition(":")
            key = k.split(";")[0].upper()
            current[key] = v
    return events[:MAX_TABLE_ROWS]


def _rows_to_table(name: str, rows: list[dict[str, Any]]) -> dict[str, Any] | None:
    if not rows:
        return None
    tables_dir().mkdir(parents=True, exist_ok=True)
    rel = Path("tables") / f"{name}.jsonl"
    path = root() / rel
    with path.open("w", encoding="utf-8") as fh:
        for row in rows[:MAX_TABLE_ROWS]:
            fh.write(json.dumps(row, default=str, ensure_ascii=False) + "\n")
    cols = list(rows[0].keys())
    return {
        "name": name,
        "path": str(rel).replace("\\", "/"),
        "columns": cols,
        "rows": min(len(rows), MAX_TABLE_ROWS),
    }


def _csv_rows(data: bytes, delimiter: str = ",") -> list[dict[str, str]]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    rows = []
    for i, row in enumerate(reader):
        if i >= MAX_TABLE_ROWS:
            break
        rows.append({str(k): ("" if v is None else str(v)) for k, v in row.items()})
    return rows


def _xlsx_rows(data: bytes) -> list[dict[str, Any]]:
    try:
        import pandas as pd

        df = pd.read_excel(io.BytesIO(data), nrows=MAX_TABLE_ROWS)
        return json.loads(df.head(MAX_TABLE_ROWS).to_json(orient="records"))
    except Exception:
        return []


def ingest_bytes(filename: str, data: bytes, *, uploaded_by: str = "chat") -> dict[str, Any]:
    if data is None:
        raise SchemaInvalidError("empty upload")
    if len(data) > MAX_BYTES:
        raise PolicyDeniedError(f"file exceeds {MAX_BYTES} bytes")
    name = safe_filename(filename)
    ext = _ext(name)
    if ext in BANNED_EXT or ext not in ALLOWED_EXT:
        raise PolicyDeniedError(f"extension '{ext or '(none)'}' is not allowed for live KB")
    resource_id = "res_" + uuid.uuid4().hex[:12]
    dest = files_dir() / f"{resource_id}_{name}"
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(data)

    kind = "text"
    table_meta = None
    text = ""
    if ext in IMAGE_EXT:
        kind = "image"
        text = f"Image attached: {name}. No OCR. Use kb_list; do not invent pixel contents."
    elif ext == ".pdf":
        kind = "document"
        text = _extract_text_pdf(data) or f"PDF stored as {name}; text extractor unavailable."
    elif ext == ".docx":
        kind = "document"
        text = _extract_text_docx(data) or f"DOCX stored as {name}; text extractor unavailable."
    elif ext == ".pptx":
        kind = "document"
        text = _extract_text_pptx(data) or f"PPTX stored as {name}; text extractor unavailable."
    elif ext == ".ics":
        kind = "calendar"
        events = _parse_ics(data.decode("utf-8", errors="replace"))
        table_meta = _rows_to_table(_unique_table(_slug(name)), events)
        text = "\n".join(f"{e.get('title')} {e.get('start_time')} {e.get('organizer')}" for e in events[:200])
    elif ext in {".csv", ".tsv"}:
        kind = "table"
        rows = _csv_rows(data, delimiter="\t" if ext == ".tsv" else ",")
        table_meta = _rows_to_table(_unique_table(_slug(name)), rows)
        text = "\n".join(json.dumps(r, ensure_ascii=False) for r in rows[:200])
    elif ext in {".xlsx", ".xls"}:
        kind = "table"
        rows = _xlsx_rows(data)
        table_meta = _rows_to_table(_unique_table(_slug(name)), rows)
        text = "\n".join(json.dumps(r, ensure_ascii=False) for r in rows[:200])
    elif ext == ".json":
        kind = "text"
        raw = data.decode("utf-8", errors="replace")
        try:
            parsed = json.loads(raw)
        except json.JSONDecodeError:
            parsed = None
        if isinstance(parsed, list) and parsed and isinstance(parsed[0], dict):
            kind = "table"
            table_meta = _rows_to_table(_unique_table(_slug(name)), parsed[:MAX_TABLE_ROWS])
        text = raw[:80_000]
    else:
        text = data.decode("utf-8", errors="replace")

    pieces = chunk_text(text) or [text[:CHUNK_SIZE] or f"(empty extract from {name})"]
    chunks = [
        {
            "id": f"{resource_id}:{i}",
            "resource_id": resource_id,
            "filename": name,
            "text": piece,
        }
        for i, piece in enumerate(pieces)
    ]
    write_chunks(resource_id, chunks)
    indexed = _try_index(chunks)
    if table_meta:
        table_meta["resource_id"] = resource_id
    record = {
        "id": resource_id,
        "filename": name,
        "kind": kind,
        "table": table_meta["name"] if table_meta else None,
        "bytes": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
        "chunks": len(chunks),
        "indexed": indexed,
        "uploaded_by": uploaded_by,
        "created_at": _now(),
        "stored_as": str(dest.relative_to(root())).replace("\\", "/"),
    }
    upsert_resource(record, table_meta)
    try:
        from retrieval.tables import clear_cache

        clear_cache()
    except Exception:
        pass
    return record
